"""Builds the consolidated output workbook in the desired output format."""
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from backend.extractor import format_date

HEADERS = ["Date", "DB Code", "DB Name", "Category", "Item Name", "MRP", "Qty in Cs"]
COLUMN_WIDTHS = [12, 10, 22, 16, 40, 10, 10]


def build_workbook(order_rows, output_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1C2B3A")
        cell.alignment = Alignment(vertical="center")

    for row in order_rows:
        ws.append(
            [
                format_date(row.order_date),
                row.db_code,
                row.db_name,
                row.category,
                row.item_name,
                row.mrp,
                row.qty,
            ]
        )

    for r in range(2, ws.max_row + 1):
        for c in range(1, 8):
            ws.cell(r, c).font = Font(name="Arial", size=10)

    for i, width in enumerate(COLUMN_WIDTHS, start=1):
        ws.column_dimensions[chr(64 + i)].width = width

    ws.freeze_panes = "A2"
    wb.save(output_path)
