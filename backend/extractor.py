"""
Parses "Terrestrial Order Format" style workbooks.

Expected layout (row numbers may vary slightly between files, so we search
for it rather than assuming fixed positions):
  - A header block containing "DB Name:-", "DB Code:-" and "Date" labels,
    each followed by its value in the next cell.
  - Further down, a table whose header row starts with "Category" and also
    contains "Item Name", "MRP", and "Order Qty in Cs".

Only item rows with a numeric, positive "Order Qty in Cs" are kept — an
MRP-only row with nothing ordered is not part of the consolidated order.
"""
from dataclasses import dataclass, field
from datetime import datetime, date

import openpyxl

MAX_HEADER_SCAN_ROWS = 20


@dataclass
class OrderRow:
    order_date: object
    db_code: object
    db_name: object
    category: object
    item_name: object
    mrp: object
    qty: float


@dataclass
class ExtractResult:
    error: str = None
    rows: list = field(default_factory=list)
    db_name: object = None
    db_code: object = None
    order_date: object = None
    sheet_name: str = None


def _col_index(header_row, name):
    target = name.strip().lower()
    for i, v in enumerate(header_row):
        if v is not None and str(v).strip().lower() == target:
            return i
    return -1


def _find_order_table(workbook):
    """Return (sheet_name, header_row_idx, header_values, all_rows) or None."""
    for sheet_name in workbook.sheetnames:
        ws = workbook[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        for r in range(min(len(rows), MAX_HEADER_SCAN_ROWS)):
            row = rows[r]
            if not row or row[0] is None:
                continue
            if str(row[0]).strip().lower() != "category":
                continue
            has_item_name = any(
                v is not None and str(v).strip().lower() == "item name" for v in row
            )
            if has_item_name:
                return sheet_name, r, row, rows
    return None


def _extract_header_info(rows, header_idx):
    db_name = db_code = order_date = None
    for r in range(header_idx):
        row = rows[r]
        if not row:
            continue
        for c, cell in enumerate(row):
            if cell is None:
                continue
            label = str(cell).strip().lower()
            if db_name is None and label.startswith("db name"):
                v = row[c + 1] if c + 1 < len(row) else None
                if v is not None and str(v).strip() != "":
                    db_name = v
            elif db_code is None and label.startswith("db code"):
                v = row[c + 1] if c + 1 < len(row) else None
                if v is not None and str(v).strip() != "":
                    db_code = v
            elif order_date is None and label.startswith("date"):
                for k in range(c + 1, len(row)):
                    if row[k] is not None and str(row[k]).strip() != "":
                        order_date = row[k]
                        break
    return db_name, db_code, order_date


def extract_from_file(file_path) -> ExtractResult:
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
    except Exception as exc:  # noqa: BLE001 - surfaced to the uploading user
        return ExtractResult(error=f"Could not open this file: {exc}")

    found = _find_order_table(wb)
    if not found:
        return ExtractResult(
            error='No "Category / Item Name" order table found in this file.'
        )
    sheet_name, header_idx, header_vals, rows = found

    cat_i = _col_index(header_vals, "Category")
    item_i = _col_index(header_vals, "Item Name")
    mrp_i = _col_index(header_vals, "MRP")
    qty_i = _col_index(header_vals, "Order Qty in Cs")

    if -1 in (cat_i, item_i, mrp_i, qty_i):
        return ExtractResult(
            error=(
                "Order table is missing an expected column "
                "(Category / Item Name / MRP / Order Qty in Cs)."
            )
        )

    db_name, db_code, order_date = _extract_header_info(rows, header_idx)

    out_rows = []
    for r in range(header_idx + 1, len(rows)):
        row = rows[r]
        if not row:
            continue
        cat = row[cat_i] if cat_i < len(row) else None
        if cat is None or str(cat).strip() == "":
            continue
        qty_raw = row[qty_i] if qty_i < len(row) else None
        try:
            qty = float(qty_raw)
        except (TypeError, ValueError):
            continue
        if not qty_raw or qty <= 0:
            continue
        out_rows.append(
            OrderRow(
                order_date=order_date,
                db_code=db_code,
                db_name=db_name,
                category=cat,
                item_name=row[item_i] if item_i < len(row) else None,
                mrp=row[mrp_i] if mrp_i < len(row) else None,
                qty=qty,
            )
        )

    return ExtractResult(
        error=None,
        rows=out_rows,
        db_name=db_name,
        db_code=db_code,
        order_date=order_date,
        sheet_name=sheet_name,
    )


def format_date(value):
    if isinstance(value, (datetime, date)):
        return value.strftime("%d-%m-%Y")
    if value is None:
        return ""
    return str(value)
